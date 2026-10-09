from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TITULO_FILL = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

CABECALHO_FILL = PatternFill(
    fill_type="solid",
    fgColor="D9EAF7"
)

ALERTA_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6"
)

MOEDA_FORMATO = 'R$ #,##0.00'
NUMERO_FORMATO = '#,##0.00'
PERCENTUAL_FORMATO = '0.00%'


def _encontrar_coluna(
    dataframe: pd.DataFrame,
    opcoes: list[str],
) -> str | None:
    """
    Localiza uma coluna ignorando maiúsculas, espaços, hífens e underscores.
    """

    colunas_normalizadas = {
        str(coluna)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", ""): coluna
        for coluna in dataframe.columns
    }

    for opcao in opcoes:
        chave = (
            opcao
            .strip()
            .lower()
            .replace(" ", "")
            .replace("_", "")
            .replace("-", "")
        )

        if chave in colunas_normalizadas:
            return colunas_normalizadas[chave]

    return None


def _primeiro_valor(
    dados: dict[str, Any] | None,
    opcoes: list[str],
    padrao: Any = 0,
) -> Any:
    """
    Procura uma métrica em um dicionário aceitando nomes alternativos.
    """

    if not dados:
        return padrao

    dados_normalizados = {
        str(chave)
        .strip()
        .lower()
        .replace(" ", "")
        .replace("_", "")
        .replace("-", ""): valor
        for chave, valor in dados.items()
    }

    for opcao in opcoes:
        chave = (
            opcao
            .strip()
            .lower()
            .replace(" ", "")
            .replace("_", "")
            .replace("-", "")
        )

        if chave in dados_normalizados:
            return dados_normalizados[chave]

    return padrao


def _numero(valor: Any, padrao: float = 0.0) -> float:
    """
    Converte valores numéricos de forma segura.
    """

    try:
        if pd.isna(valor):
            return padrao

        return float(valor)
    except (TypeError, ValueError):
        return padrao


def _formatar_moeda(valor: Any) -> str:
    """
    Formata valores monetários no padrão brasileiro para o resumo.
    """

    valor_numerico = _numero(valor)

    texto = f"{valor_numerico:,.2f}"
    texto = texto.replace(",", "X").replace(".", ",").replace("X", ".")

    return f"R$ {texto}"


def _formatar_percentual(valor: Any) -> str:
    """
    Aceita 0.15 ou 15 como 15,00%.
    """

    valor_numerico = _numero(valor)

    if valor_numerico > 1:
        valor_numerico = valor_numerico / 100

    return f"{valor_numerico:.2%}".replace(".", ",")


def _padronizar_dataframe(
    dataframe: pd.DataFrame | None,
) -> pd.DataFrame:
    """
    Garante que um DataFrame vazio seja retornado em vez de None.
    """

    if dataframe is None:
        return pd.DataFrame()

    return dataframe.copy()


def _criar_resumo(
    dados: pd.DataFrame,
    metricas: dict[str, Any] | None,
    economia_potencial: float | int | None,
) -> pd.DataFrame:
    """
    Monta a aba de resumo executivo usando métricas recebidas
    ou calculando valores diretamente do DataFrame.
    """

    dados = _padronizar_dataframe(dados)
    metricas = metricas or {}

    coluna_valor = _encontrar_coluna(
        dados,
        ["valor", "valor_total", "custo", "custo_total", "valor_rs"]
    )

    coluna_veiculo = _encontrar_coluna(
        dados,
        ["veiculo", "placa", "id_veiculo", "idveiculo"]
    )

    coluna_data = _encontrar_coluna(
        dados,
        ["data", "data_lancamento", "data_transacao", "dt_lancamento"]
    )

    custo_calculado = (
        pd.to_numeric(dados[coluna_valor], errors="coerce").fillna(0).sum()
        if coluna_valor
        else 0
    )

    veiculos_calculados = (
        dados[coluna_veiculo].nunique()
        if coluna_veiculo
        else 0
    )

    transacoes_calculadas = len(dados)

    custo_total = _primeiro_valor(
        metricas,
        [
            "custo_total",
            "total_custos",
            "custo_total_frota",
            "valor_total"
        ],
        custo_calculado
    )

    total_veiculos = _primeiro_valor(
        metricas,
        [
            "total_veiculos",
            "quantidade_veiculos",
            "veiculos_analisados"
        ],
        veiculos_calculados
    )

    total_transacoes = _primeiro_valor(
        metricas,
        [
            "total_transacoes",
            "quantidade_transacoes",
            "transacoes_analisadas"
        ],
        transacoes_calculadas
    )

    custo_medio_veiculo = _primeiro_valor(
        metricas,
        [
            "custo_medio_veiculo",
            "custo_medio_por_veiculo"
        ],
        _numero(custo_total) / total_veiculos if total_veiculos else 0
    )

    economia = _numero(economia_potencial)

    if coluna_data and not dados.empty:
        datas = pd.to_datetime(
            dados[coluna_data],
            errors="coerce"
        ).dropna()

        if not datas.empty:
            periodo = (
                f"{datas.min().strftime('%d/%m/%Y')} até "
                f"{datas.max().strftime('%d/%m/%Y')}"
            )
        else:
            periodo = "Não identificado"
    else:
        periodo = "Não identificado"

    percentual_economia = (
        economia / _numero(custo_total)
        if _numero(custo_total) > 0
        else 0
    )

    resumo = pd.DataFrame(
        {
            "Indicador": [
                "Período analisado",
                "Total de veículos",
                "Total de transações",
                "Custo total da frota",
                "Custo médio por veículo",
                "Economia potencial estimada",
                "Potencial de redução de custos",
                "Data de geração do relatório"
            ],
            "Valor": [
                periodo,
                int(_numero(total_veiculos)),
                int(_numero(total_transacoes)),
                _formatar_moeda(custo_total),
                _formatar_moeda(custo_medio_veiculo),
                _formatar_moeda(economia),
                _formatar_percentual(percentual_economia),
                datetime.now().strftime("%d/%m/%Y %H:%M")
            ]
        }
    )

    return resumo


def _criar_custos_categoria(
    dados: pd.DataFrame,
) -> pd.DataFrame:
    """
    Cria o consolidado financeiro por categoria.
    """

    dados = _padronizar_dataframe(dados)

    coluna_categoria = _encontrar_coluna(
        dados,
        ["categoria", "tipo_despesa", "tipo", "grupo"]
    )

    coluna_valor = _encontrar_coluna(
        dados,
        ["valor", "valor_total", "custo", "custo_total", "valor_rs"]
    )

    if not coluna_categoria or not coluna_valor:
        return pd.DataFrame(
            {
                "Mensagem": [
                    "Não foi possível identificar as colunas "
                    "de categoria e valor para consolidar os custos."
                ]
            }
        )

    dados_analise = dados.copy()
    dados_analise[coluna_valor] = pd.to_numeric(
        dados_analise[coluna_valor],
        errors="coerce"
    ).fillna(0)

    custos = (
        dados_analise
        .groupby(coluna_categoria, dropna=False)
        .agg(
            **{
                "Custo total (R$)": (
                    coluna_valor,
                    "sum"
                ),
                "Custo médio (R$)": (
                    coluna_valor,
                    "mean"
                ),
                "Quantidade de transações": (
                    coluna_valor,
                    "count"
                )
            }
        )
        .reset_index()
        .rename(
            columns={
                coluna_categoria: "Categoria"
            }
        )
        .sort_values(
            by="Custo total (R$)",
            ascending=False
        )
    )

    custo_geral = custos["Custo total (R$)"].sum()

    custos["Participação no custo total"] = (
        custos["Custo total (R$)"] / custo_geral
        if custo_geral > 0
        else 0
    )

    return custos


def _criar_veiculos_criticos(
    dados: pd.DataFrame,
    benchmark: pd.DataFrame | None,
) -> pd.DataFrame:
    """
    Consolida custos por veículo e incorpora colunas relevantes
    do benchmark, quando disponível.
    """

    dados = _padronizar_dataframe(dados)
    benchmark = _padronizar_dataframe(benchmark)

    coluna_veiculo = _encontrar_coluna(
        dados,
        ["veiculo", "placa", "id_veiculo", "idveiculo"]
    )

    coluna_valor = _encontrar_coluna(
        dados,
        ["valor", "valor_total", "custo", "custo_total", "valor_rs"]
    )

    if not coluna_veiculo or not coluna_valor:
        return pd.DataFrame(
            {
                "Mensagem": [
                    "Não foi possível identificar as colunas "
                    "de veículo e valor para gerar o ranking."
                ]
            }
        )

    dados_analise = dados.copy()
    dados_analise[coluna_valor] = pd.to_numeric(
        dados_analise[coluna_valor],
        errors="coerce"
    ).fillna(0)

    veiculos = (
        dados_analise
        .groupby(coluna_veiculo, dropna=False)
        .agg(
            **{
                "Custo total (R$)": (
                    coluna_valor,
                    "sum"
                ),
                "Quantidade de transações": (
                    coluna_valor,
                    "count"
                )
            }
        )
        .reset_index()
        .rename(
            columns={
                coluna_veiculo: "Veículo"
            }
        )
        .sort_values(
            by="Custo total (R$)",
            ascending=False
        )
    )

    if not benchmark.empty:
        coluna_veiculo_benchmark = _encontrar_coluna(
            benchmark,
            ["veiculo", "placa", "id_veiculo", "idveiculo"]
        )

        if coluna_veiculo_benchmark:
            benchmark_ajustado = benchmark.copy()

            if coluna_veiculo_benchmark != "Veículo":
                benchmark_ajustado = benchmark_ajustado.rename(
                    columns={
                        coluna_veiculo_benchmark: "Veículo"
                    }
                )

            colunas_importantes = [
                coluna
                for coluna in benchmark_ajustado.columns
                if coluna != "Veículo"
            ]

            veiculos = veiculos.merge(
                benchmark_ajustado[
                    ["Veículo"] + colunas_importantes
                ],
                on="Veículo",
                how="left"
            )

    return veiculos.head(20)


def _criar_anomalias(
    anomalias: pd.DataFrame | None,
) -> pd.DataFrame:
    """
    Seleciona e ordena anomalias sem exigir um formato único.
    """

    anomalias = _padronizar_dataframe(anomalias)

    if anomalias.empty:
        return pd.DataFrame(
            {
                "Mensagem": [
                    "Nenhuma anomalia foi identificada com os critérios atuais."
                ]
            }
        )

    coluna_prioridade = _encontrar_coluna(
        anomalias,
        ["prioridade", "nivel_risco", "risco", "severidade"]
    )

    coluna_valor = _encontrar_coluna(
        anomalias,
        ["valor", "valor_total", "custo", "custo_total", "valor_rs"]
    )

    ordem = anomalias.copy()

    if coluna_prioridade:
        mapa_prioridade = {
            "CRÍTICA": 1,
            "CRITICA": 1,
            "ALTA": 2,
            "MÉDIA": 3,
            "MEDIA": 3,
            "BAIXA": 4
        }

        ordem["_ordem_prioridade"] = (
            ordem[coluna_prioridade]
            .astype(str)
            .str.upper()
            .map(mapa_prioridade)
            .fillna(5)
        )

        ordem = ordem.sort_values(
            by="_ordem_prioridade",
            ascending=True
        )

        ordem = ordem.drop(
            columns="_ordem_prioridade"
        )

    if coluna_valor:
        ordem[coluna_valor] = pd.to_numeric(
            ordem[coluna_valor],
            errors="coerce"
        )

        ordem = ordem.sort_values(
            by=coluna_valor,
            ascending=False
        )

    return ordem.head(200)


def _criar_plano_acao(
    recomendacoes: pd.DataFrame | list[dict[str, Any]] | None,
) -> pd.DataFrame:
    """
    Converte recomendações em um DataFrame adequado para a aba de ação.
    """

    if recomendacoes is None:
        return pd.DataFrame(
            {
                "Mensagem": [
                    "Nenhuma recomendação foi gerada."
                ]
            }
        )

    if isinstance(recomendacoes, list):
        plano = pd.DataFrame(recomendacoes)
    elif isinstance(recomendacoes, pd.DataFrame):
        plano = recomendacoes.copy()
    else:
        return pd.DataFrame(
            {
                "Mensagem": [
                    "Formato de recomendações não reconhecido."
                ]
            }
        )

    if plano.empty:
        return pd.DataFrame(
            {
                "Mensagem": [
                    "Nenhuma recomendação foi gerada."
                ]
            }
        )

    coluna_prioridade = _encontrar_coluna(
        plano,
        ["prioridade", "nivel_prioridade", "severidade"]
    )

    if coluna_prioridade:
        mapa_prioridade = {
            "CRÍTICA": 1,
            "CRITICA": 1,
            "ALTA": 2,
            "MÉDIA": 3,
            "MEDIA": 3,
            "BAIXA": 4
        }

        plano["_ordem_prioridade"] = (
            plano[coluna_prioridade]
            .astype(str)
            .str.upper()
            .map(mapa_prioridade)
            .fillna(5)
        )

        plano = plano.sort_values(
            by="_ordem_prioridade",
            ascending=True
        )

        plano = plano.drop(
            columns="_ordem_prioridade"
        )

    return plano


def _estilizar_planilha(
    planilha,
    dataframe: pd.DataFrame,
    linha_cabecalho: int = 1,
) -> None:
    """
    Aplica estilos básicos e ajusta a largura de colunas no Excel.
    """

    planilha.freeze_panes = f"A{linha_cabecalho + 1}"
    planilha.auto_filter.ref = planilha.dimensions

    for celula in planilha[linha_cabecalho]:
        celula.fill = CABECALHO_FILL
        celula.font = Font(
            bold=True,
            color="000000"
        )
        celula.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    for coluna in range(1, planilha.max_column + 1):
        letra = get_column_letter(coluna)
        maior_tamanho = 0

        for celula in planilha[letra]:
            valor = "" if celula.value is None else str(celula.value)
            maior_tamanho = max(
                maior_tamanho,
                len(valor)
            )

        planilha.column_dimensions[letra].width = min(
            max(maior_tamanho + 2, 12),
            45
        )

    for linha in planilha.iter_rows(
        min_row=linha_cabecalho + 1
    ):
        for celula in linha:
            celula.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    for indice, coluna in enumerate(
        dataframe.columns,
        start=1
    ):
        nome_coluna = str(coluna).lower()

        if (
            "custo" in nome_coluna
            or "valor" in nome_coluna
            or "economia" in nome_coluna
            or "impacto" in nome_coluna
        ):
            for linha in range(
                linha_cabecalho + 1,
                planilha.max_row + 1
            ):
                planilha.cell(
                    row=linha,
                    column=indice
                ).number_format = MOEDA_FORMATO

        if (
            "participação" in nome_coluna
            or "participacao" in nome_coluna
            or "percentual" in nome_coluna
            or "%" in nome_coluna
        ):
            for linha in range(
                linha_cabecalho + 1,
                planilha.max_row + 1
            ):
                planilha.cell(
                    row=linha,
                    column=indice
                ).number_format = PERCENTUAL_FORMATO


def _escrever_aba(
    writer: pd.ExcelWriter,
    nome_aba: str,
    dataframe: pd.DataFrame,
    subtitulo: str | None = None,
) -> None:
    """
    Escreve título, subtítulo e DataFrame em uma aba formatada.
    """

    dataframe = _padronizar_dataframe(dataframe)

    if dataframe.empty:
        dataframe = pd.DataFrame(
            {
                "Mensagem": [
                    "Não há dados disponíveis para esta seção."
                ]
            }
        )

    linha_inicial = 3 if subtitulo else 2

    dataframe.to_excel(
        writer,
        sheet_name=nome_aba,
        index=False,
        startrow=linha_inicial
    )

    planilha = writer.sheets[nome_aba]

    planilha.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=max(len(dataframe.columns), 1)
    )

    celula_titulo = planilha.cell(
        row=1,
        column=1
    )

    celula_titulo.value = f"FleetWise — {nome_aba}"
    celula_titulo.fill = TITULO_FILL
    celula_titulo.font = Font(
        bold=True,
        color="FFFFFF",
        size=14
    )
    celula_titulo.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    planilha.row_dimensions[1].height = 25

    if subtitulo:
        planilha.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=max(len(dataframe.columns), 1)
        )

        celula_subtitulo = planilha.cell(
            row=2,
            column=1
        )

        celula_subtitulo.value = subtitulo
        celula_subtitulo.font = Font(
            italic=True,
            color="666666"
        )
        celula_subtitulo.alignment = Alignment(
            wrap_text=True
        )

    _estilizar_planilha(
        planilha,
        dataframe,
        linha_cabecalho=linha_inicial + 1
    )


def gerar_relatorio_excel(
    dados: pd.DataFrame,
    metricas: dict[str, Any] | None = None,
    anomalias: pd.DataFrame | None = None,
    benchmark: pd.DataFrame | None = None,
    recomendacoes: pd.DataFrame | list[dict[str, Any]] | None = None,
    economia_potencial: float | int | None = 0,
) -> bytes:
    """
    Gera o relatório executivo FleetWise em Excel.

    Parâmetros
    ----------
    dados:
        DataFrame principal após ETL e filtros do dashboard.
    metricas:
        Dicionário opcional com KPIs já calculados.
    anomalias:
        DataFrame com lançamentos ou veículos classificados como anômalos.
    benchmark:
        DataFrame com indicadores de benchmark por veículo.
    recomendacoes:
        DataFrame ou lista de dicionários com o plano de ação.
    economia_potencial:
        Valor total estimado de economia.

    Retorno
    -------
    bytes:
        Conteúdo do arquivo .xlsx, pronto para ser usado no
        st.download_button do Streamlit.
    """

    resumo = _criar_resumo(
        dados=dados,
        metricas=metricas,
        economia_potencial=economia_potencial
    )

    custos_categoria = _criar_custos_categoria(
        dados=dados
    )

    veiculos_criticos = _criar_veiculos_criticos(
        dados=dados,
        benchmark=benchmark
    )

    tabela_anomalias = _criar_anomalias(
        anomalias=anomalias
    )

    plano_acao = _criar_plano_acao(
        recomendacoes=recomendacoes
    )

    buffer = BytesIO()

    with pd.ExcelWriter(
        buffer,
        engine="openpyxl"
    ) as writer:
        _escrever_aba(
            writer=writer,
            nome_aba="Resumo Executivo",
            dataframe=resumo,
            subtitulo=(
                "Visão consolidada dos custos, indicadores e "
                "oportunidades identificadas pela análise FleetWise."
            )
        )

        _escrever_aba(
            writer=writer,
            nome_aba="Custos por Categoria",
            dataframe=custos_categoria,
            subtitulo=(
                "Consolidação financeira das despesas por categoria."
            )
        )

        _escrever_aba(
            writer=writer,
            nome_aba="Veículos Críticos",
            dataframe=veiculos_criticos,
            subtitulo=(
                "Ranking dos veículos com maior impacto financeiro "
                "e dados de benchmark quando disponíveis."
            )
        )

        _escrever_aba(
            writer=writer,
            nome_aba="Anomalias",
            dataframe=tabela_anomalias,
            subtitulo=(
                "Registros potencialmente atípicos que exigem "
                "validação operacional."
            )
        )

        _escrever_aba(
            writer=writer,
            nome_aba="Plano de Ação",
            dataframe=plano_acao,
            subtitulo=(
                "Recomendações priorizadas para redução de custos "
                "e melhoria da eficiência da frota."
            )
        )

    buffer.seek(0)

    return buffer.getvalue()